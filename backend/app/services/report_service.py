import csv
import io
from datetime import datetime, timedelta
from sqlalchemy import select, or_, cast, Text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.vessel import AisVesselStatic, VesselState
from app.models.reference import AisShipType, AisNavStatus
from app.services.encounter_service import detect_encounters

MAX_EXPORT_ROWS = 10000


async def fetch_vessels_rows(
    db: AsyncSession, search: str | None, ship_type: str | None
) -> tuple[list[str], list[list]]:
    base = (
        select(
            AisVesselStatic.mmsi,
            AisVesselStatic.name,
            AisVesselStatic.call_sign,
            AisVesselStatic.imo,
            AisShipType.group_name.label("ship_type_group"),
            AisVesselStatic.length_m,
            AisVesselStatic.width_m,
            AisVesselStatic.draught_m,
            AisNavStatus.name.label("last_nav_status"),
            VesselState.last_sog,
            VesselState.last_received_at,
        )
        .outerjoin(AisShipType, AisVesselStatic.ship_type_code == AisShipType.code)
        .outerjoin(VesselState, AisVesselStatic.mmsi == VesselState.mmsi)
        .outerjoin(AisNavStatus, VesselState.last_nav_status == AisNavStatus.code)
    )

    if search:
        term = f"%{search}%"
        base = base.where(
            or_(
                AisVesselStatic.name.ilike(term),
                AisVesselStatic.call_sign.ilike(term),
                cast(AisVesselStatic.mmsi, Text).ilike(term),
            )
        )
    if ship_type:
        base = base.where(AisShipType.group_name == ship_type)

    rows = (await db.execute(base.order_by(AisVesselStatic.name).limit(MAX_EXPORT_ROWS))).all()

    columns = [
        "MMSI", "Nama", "Call Sign", "IMO", "Tipe", "Panjang (m)",
        "Lebar (m)", "Sarat (m)", "Status", "SOG (kn)", "Update Terakhir",
    ]
    data = [
        [
            r.mmsi, r.name, r.call_sign, r.imo, r.ship_type_group,
            r.length_m, r.width_m, r.draught_m, r.last_nav_status, r.last_sog,
            r.last_received_at.strftime("%Y-%m-%d %H:%M") if r.last_received_at else "",
        ]
        for r in rows
    ]
    return columns, data


async def fetch_encounters_rows(
    db: AsyncSession,
    date_from: datetime | None,
    date_to: datetime | None,
    distance_m: float,
) -> tuple[list[str], list[list]]:
    items, _, _ = await detect_encounters(db, date_from, date_to, distance_m, limit=1000)

    columns = [
        "Kapal A", "MMSI A", "Kapal B", "MMSI B", "Mulai", "Selesai",
        "Durasi (menit)", "Jarak Min (m)", "Jumlah Sampel",
    ]
    data = [
        [
            it.name_a or "-", it.mmsi_a, it.name_b or "-", it.mmsi_b,
            it.first_seen.strftime("%Y-%m-%d %H:%M"), it.last_seen.strftime("%Y-%m-%d %H:%M"),
            it.duration_minutes, it.min_distance_m, it.sample_count,
        ]
        for it in items
    ]
    return columns, data


def build_csv(columns: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def build_excel(columns: list[str], rows: list[list], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(columns)
    header_fill = PatternFill(start_color="1A4E8F", end_color="1A4E8F", fill_type="solid")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for row in rows:
        ws.append(row)
    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(columns: list[str], rows: list[list], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(title, styles["Title"]),
        Paragraph(f"Diekspor: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
        Spacer(1, 12),
    ]

    data = [columns] + [["" if c is None else str(c) for c in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1A4E8F")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f0f0")]),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()
